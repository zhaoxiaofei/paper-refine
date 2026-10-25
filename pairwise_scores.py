#!/usr/bin/env python3
"""Build a pairwise-score workbook from manuscript-judging CSV files.

Each worksheet is one round. The tab is the run directory with the shared
``cnb-`` prefix removed and ``_rN`` appended, for example
``12to13-092216-e1cdf83_r1``.

Rows and columns are version ids from ``member`` and ``opponent_id``
(``orig``, ``a2``, ``i1``, ``w1``, …), not the ``v1`` / ``v2`` labels.
Row 1 is the column names and column A is the row names. The cell at
row version R and column version C is the mean ``credited`` score of R
against C. Positive means R was judged better than C. A versus B is not
the negative of B versus A: the opposite direction is a separate judgment
and is left blank when it was not judged. By default only the pipeline's
``direction=own`` rows are used; each ``direction=received`` row is the
negation of a judgment already counted in the opponent's own frame, not a
second judgment. The first sheet, Description, explains the workbook.

Usage
-----
    python pairwise_scores.py /path/to/cnb-manuscript-postgrad-data -o pairwise_scores.xlsx
    python pairwise_scores.py round1_raw_scores.csv round2_raw_scores.csv -o out.xlsx
    python pairwise_scores.py head-dump.txt -o out.xlsx

A text file produced by ``head **/reports/round*.csv`` (``==> path <==``
markers) is accepted, but only the lines present in that dump are scored.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path

csv.field_size_limit(min(sys.maxsize, 10_000_000))

# `paper_pipeline.py` writes BOTH frames of one judgment: the target's
# `direction=own` row and the opponent's `direction=received` negation. Keeping
# the latter in a directed matrix silently manufactures a judgment in the
# opposite direction -- exactly what this workbook says it never does.
HEAD_MARKER = re.compile(r"^==>\s+(.*?)\s+<==\s*$")
ROUND_NAME = re.compile(r"round(\d+)", re.IGNORECASE)


def version_sort_key(label: str) -> tuple:
    """``orig`` first, then ids such as ``a2``, ``i1``, ``r1_i3``, ``w2``."""
    text = str(label).strip()
    lowered = text.lower()
    if lowered in {"orig", "original", "target"}:
        return (0, ())
    chunks = re.findall(r"\d+|[A-Za-z]+", lowered)
    key: list[tuple] = []
    for chunk in chunks:
        if chunk.isdigit():
            key.append((0, int(chunk)))
        else:
            key.append((1, chunk))
    return (1, tuple(key), lowered)


def run_name_for(csv_path: Path) -> str:
    """Sheet name: parent of ``reports/``, else the CSV's parent directory."""
    parent = csv_path.parent
    if parent.name == "reports":
        name = parent.parent.name
    else:
        name = parent.name
    if name in {"", ".", "/"}:
        name = csv_path.stem
    return name


def round_from_filename(path: Path) -> str | None:
    match = ROUND_NAME.search(path.name)
    return match.group(1) if match else None


def parse_score(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        score = float(text)
    except ValueError:
        return None
    return score if math.isfinite(score) else None


def is_head_dump(path: Path) -> bool:
    if path.suffix.lower() == ".csv":
        return False
    try:
        with path.open("r", encoding="utf-8", errors="replace") as handle:
            chunk = handle.read(8000)
    except OSError:
        return False
    return "==>" in chunk and "opponent_label" in chunk


def iter_head_dump(path: Path):
    """Yield (virtual csv path, text) blocks from a ``head`` listing."""
    current: str | None = None
    buffer: list[str] = []

    def flush():
        if current and buffer:
            yield Path(current), "\n".join(buffer) + "\n"

    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for raw in handle:
            line = raw.rstrip("\n")
            marker = HEAD_MARKER.match(line.strip())
            if marker:
                yield from flush()
                current = marker.group(1)
                buffer = []
                continue
            if current is None:
                continue
            if not line.strip():
                continue
            # Drop a trailing shell prompt that ``head`` sessions sometimes append.
            # A CSV row can legitimately end in "$" (for example in a free-text
            # reason column); prompt lines carry no comma-separated fields.
            if "," not in line and (line.startswith("(base) ") or line.endswith("$")):
                continue
            buffer.append(line)
    yield from flush()


def discover_csv_files(inputs: list[Path]) -> list[tuple[Path, str | None]]:
    """Return (path, optional inline text) pairs.

    Inline text is set only for blocks reconstructed from a head dump.
    """
    found: list[tuple[Path, str | None]] = []
    for item in inputs:
        if not item.exists():
            raise SystemExit(f"input not found: {item}")
        if item.is_file() and is_head_dump(item):
            blocks = list(iter_head_dump(item))
            if not blocks:
                raise SystemExit(f"no CSV blocks found in head dump: {item}")
            found.extend((path, text) for path, text in blocks)
            continue
        if item.is_file():
            found.append((item, None))
            continue
        csvs = sorted(
            p for p in item.rglob("*") if p.is_file() and p.suffix.lower() == ".csv"
        )
        preferred = [
            p
            for p in csvs
            if re.search(r"round\d+_raw_scores\.csv$", p.name, re.IGNORECASE)
        ]
        chosen = preferred or csvs
        if not chosen:
            raise SystemExit(f"no CSV files under {item}")
        found.extend((p, None) for p in chosen)
    if not found:
        raise SystemExit("no input CSV files")
    return found


def read_rows(path: Path, inline: str | None) -> list[dict[str, str]]:
    if inline is None:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    return list(csv.DictReader(inline.splitlines()))


def opponent_of(row: dict[str, str]) -> str:
    """Version id of the opponent. ``v1``/``v2`` labels are not version names."""
    return (row.get("opponent_id") or "").strip()


def aggregate(values: list[float], how: str) -> float:
    if how == "median":
        return float(statistics.median(values))
    if how == "sum":
        return float(sum(values))
    return float(statistics.fmean(values))


class PairScores:
    def __init__(self) -> None:
        self.values: list[float] = []
        self.judges: list[str] = []

    def add(self, score: float, judge: str) -> None:
        self.values.append(score)
        self.judges.append(judge)


def collect(inputs: list[Path], how: str, pool_rounds: bool, own_only: bool):
    """Return run -> round -> {(member, opponent): PairScores} plus notes."""
    files = discover_csv_files(inputs)
    # run -> round -> pair -> PairScores
    cubes: dict[str, dict[str, dict[tuple[str, str], PairScores]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(PairScores))
    )
    sources: dict[str, list[str]] = defaultdict(list)
    seen: set[tuple] = set()
    directions_seen: set[str] = set()
    skipped_direction = 0
    skipped_incomplete = 0
    truncated = False

    for path, inline in files:
        if inline is not None:
            truncated = True
        run = run_name_for(path)
        file_round = round_from_filename(path)
        try:
            rows = read_rows(path, inline)
        except OSError as exc:
            raise SystemExit(f"could not read {path}: {exc}") from exc
        data_rows = 0
        for row in rows:
            member = (row.get("member") or "").strip()
            opponent = opponent_of(row)
            score = parse_score(row.get("credited"))
            score_from_written = False
            if score is None:
                score = parse_score(row.get("score_as_written"))
                score_from_written = True
            direction = (row.get("direction") or "").strip()
            if direction:
                directions_seen.add(direction)
            if own_only and direction and direction.lower() != "own":
                skipped_direction += 1
                continue
            if not member or not opponent or score is None or member == opponent:
                skipped_incomplete += 1
                continue
            # A received row is the opponent's frame: when the caller asks to
            # keep both directions for audit, its credited value is the
            # NEGATION of the judgment, never a second identical observation.
            if direction.lower() == "received" and score_from_written:
                # A received row with no credited value carries the opponent's
                # frame in `score_as_written`; the credited value is absent
                # because the direction is derived. A received row that already
                # carries `credited` (the normal pipeline twin) is already the
                # negation and is kept verbatim for the audit.
                score = -score
            source_round = (row.get("round") or "").strip() or file_round or "1"
            round_id = source_round
            if pool_rounds:
                round_id = "all"
            judge = (row.get("source_judge_run") or "").strip()
            judge_index = (row.get("source_judge_index") or "").strip()
            dedupe_key = (
                run,
                source_round,
                round_id,
                member,
                opponent,
                judge,
                judge_index,
                score,
                (row.get("sheet_target") or "").strip(),
            )
            # A bare judge identity ("", "") is not an identity: two blank-
            # identity rows can be two independent judgments, so they are kept.
            if dedupe_key in seen and (judge or judge_index):
                continue
            seen.add(dedupe_key)
            # Keep this direction only. Never store (opponent, member) = -score.
            cubes[run][round_id][(member, opponent)].add(
                score, judge or judge_index or "judge"
            )
            data_rows += 1
        label = path.name
        if file_round and f"round{file_round}" not in path.name.lower():
            label = f"{path.name} (round {file_round})"
        sources[run].append(f"{label}: {data_rows} score row{'s' if data_rows != 1 else ''}")

    notes = {
        "truncated": truncated,
        "sources": sources,
        "directions": sorted(directions_seen),
        "skipped_direction": skipped_direction,
        "skipped_incomplete": skipped_incomplete,
        "agg": how,
        "pooled": pool_rounds,
    }
    return cubes, notes


def shared_name_prefix(names: list[str]) -> str:
    """Longest shared prefix that ends on ``-`` or ``_`` (never a partial token)."""
    if len(names) < 2:
        return ""
    prefix = names[0]
    for name in names[1:]:
        limit = min(len(prefix), len(name))
        index = 0
        while index < limit and prefix[index] == name[index]:
            index += 1
        prefix = prefix[:index]
        if not prefix:
            return ""
    cut = max(prefix.rfind("-"), prefix.rfind("_"))
    if cut < 0:
        return ""
    return prefix[: cut + 1]


def sheet_stem(run: str, prefix: str) -> str:
    """Drop a leading ``cnb-`` / ``cnb_`` and any other shared prefix."""
    original = run.strip()
    text = original
    for token in ("cnb-", "cnb_"):
        if text.lower().startswith(token):
            text = text[len(token) :]
            break
    # Callers may compute the shared prefix before or after dropping ``cnb-``.
    # Accept both: strip it from the raw name when the raw name carries it.
    if prefix and original.startswith(prefix):
        text = original[len(prefix) :]
    if prefix and text.startswith(prefix):
        text = text[len(prefix) :]
    return text.strip("-_") or original


def sheet_tab(stem: str, round_id: str, used: set[str]) -> str:
    """``{stem}_rN``, shortened only if Excel's 31-character limit requires it."""
    if str(round_id).isdigit():
        suffix = f"_r{int(round_id)}"
    elif round_id == "all":
        suffix = "_all"
    else:
        safe_round = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(round_id))
        suffix = f"_{safe_round}"
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", stem).strip() or "run"
    if len(suffix) > 31:
        suffix = suffix[:31]
    room = 31 - len(suffix)
    base = cleaned if len(cleaned) + len(suffix) <= 31 else cleaned[:room]
    candidate = (base + suffix)[:31]
    index = 2
    while candidate in used:
        tag = f"~{index}"
        base = cleaned[: max(room - len(tag), 1)]
        candidate = (base + tag + suffix)[:31]
        index += 1
    used.add(candidate)
    return candidate


def directed_score_note(cubes, how: str) -> str:
    """Describe whether any pair was judged in both directions."""
    both = 0
    exact_negations = 0
    examples: list[str] = []
    for rounds in cubes.values():
        for pairs in rounds.values():
            seen_pairs: set[tuple[str, str]] = set()
            for member, opponent in pairs:
                if (opponent, member) not in pairs or (opponent, member) in seen_pairs:
                    continue
                seen_pairs.add((member, opponent))
                both += 1
                left = aggregate(pairs[(member, opponent)].values, how)
                right = aggregate(pairs[(opponent, member)].values, how)
                if abs(left + right) < 1e-9:
                    exact_negations += 1
                elif len(examples) < 3:
                    examples.append(
                        f"{member} vs {opponent} is {left:g}, while {opponent} vs {member} is {right:g}"
                    )
    if both == 0:
        return (
            "In this file, no pair was judged in both directions. "
            "The missing direction is left blank. It is not filled with the negated score."
        )
    pair_word = "pair" if both == 1 else "pairs"
    negation_word = "is an exact negation" if exact_negations == 1 else "are exact negations"
    verb = "was" if both == 1 else "were"
    text = (
        f"{both} {pair_word} {verb} judged in both directions. "
        f"{exact_negations} of those {negation_word}; the others are not."
    )
    if examples:
        text += " For example: " + "; ".join(examples) + "."
    return text


def add_description_sheet(workbook, cubes, index_rows, how: str) -> None:
    """First sheet: how the score sheets are built and how to read a cell."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet = workbook.create_sheet("Description", 0)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 120
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    title_font = Font(name="Calibri", bold=True, size=16, color="1A1A1A")
    section_font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    label_font = Font(name="Calibri", bold=True, size=11, color="1A1A1A")
    body_font = Font(name="Calibri", size=11, color="1A1A1A")
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    agg_word = {"mean": "mean", "median": "median", "sum": "sum"}[how]

    sheet["A1"] = "How to read this workbook"
    sheet["A1"].font = title_font
    sheet.merge_cells("A1:B1")
    sheet.row_dimensions[1].height = 24

    sheet["A3"] = "Score sheets"
    sheet["A3"].font = section_font

    rules = [
        ("Sheet name", "The run directory, without a shared leading prefix such as cnb-, then _rN. Example: 12to13-092216-e1cdf83_r1 is round 1 of cnb-12to13-092216-e1cdf83."),
        ("Row 1", "Column names. These are the opponent version ids."),
        ("Column A", "Row names. These are the version ids the score is credited to."),
        ("Scores", f"The cell in row R and column C is the {agg_word} of the raw credited values for that ordered pair. If more than one judge scored the pair, the cell comment lists each judge."),
        ("Sign", "Positive means the row version was judged better than the column version. Negative means the row version was judged worse. Zero means the judges called them equal."),
        ("Not symmetric", "A versus B is not the negative of B versus A. Each direction is its own judgment and can differ in size and in sign. If one direction was not judged, that cell is blank. The blank is never filled by flipping the sign of the other direction."),
        ("Diagonal", "An em dash. A version is not compared with itself."),
        ("Version ids", "Taken from member and opponent_id: orig, a2, i1, i2, w1, w2, r1_i3, and similar ids. The v1, v2, v3 values in opponent_label are not version names and are not used as row or column names."),
        ("This file", directed_score_note(cubes, how)),
        ("Extra", "The output associated with the current sheet is given as the input to the round associated with the next sheet. If the next round is characterized by the suffix `_r1`, then another instance of paper_pipeline.py was launched to start the next round using the champion of the current run.")
    ]

    header_row = 4
    for col, text in ((1, "Item"), (2, "Meaning")):
        cell = sheet.cell(header_row, col, text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin
    for offset, (item, meaning) in enumerate(rules, start=1):
        row = header_row + offset
        item_cell = sheet.cell(row, 1, item)
        item_cell.font = label_font
        item_cell.alignment = wrap
        item_cell.border = thin
        meaning_cell = sheet.cell(row, 2, meaning)
        meaning_cell.font = body_font
        meaning_cell.alignment = wrap
        meaning_cell.border = thin
        sheet.row_dimensions[row].height = 48 if item != "This file" else 36

    index_title = header_row + len(rules) + 2
    sheet.cell(index_title, 1, "Sheets in this file").font = section_font
    columns = ("Sheet", "Directory", "Round", "Versions", "Directed pairs")
    for col, text in enumerate(columns, start=1):
        cell = sheet.cell(index_title + 1, col, text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin
    for offset, (tab, run, round_id, n_versions, n_pairs) in enumerate(index_rows):
        row = index_title + 2 + offset
        values = (
            tab,
            run,
            f"r{round_id}" if str(round_id).isdigit() else str(round_id),
            n_versions,
            n_pairs,
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row, col, value)
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin
        sheet.row_dimensions[row].height = 18

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 110
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 18
    sheet.auto_filter.ref = None
    sheet.freeze_panes = "A5"
    sheet.oddHeader.left.text = "Description"


def write_workbook(cubes, notes, output: Path, how: str) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    label_fill = PatternFill("solid", fgColor="1F4E79")
    label_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    score_font = Font(name="Calibri", size=11)
    diagonal_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    used_names: set[str] = set()
    # Compute the shared prefix AFTER the cnb- prefix is removed, because
    # sheet_stem() removes cnb- before it tries to remove the shared prefix.
    name_prefix = shared_name_prefix(sorted(sheet_stem(run, "") for run in cubes))

    def round_sort_key(label: str):
        return (0, int(label)) if str(label).isdigit() else (1, str(label))

    agg_word = {"mean": "mean", "median": "median", "sum": "sum"}[how]
    index_rows: list[tuple] = []
    for run in sorted(cubes):
        stem = sheet_stem(run, name_prefix)
        for round_id in sorted(cubes[run], key=round_sort_key):
            pairs = cubes[run][round_id]
            versions = sorted({a for pair in pairs for a in pair}, key=version_sort_key)
            tab_name = sheet_tab(stem, round_id, used_names)
            index_rows.append((tab_name, run, round_id, len(versions), len(pairs)))
            sheet = workbook.create_sheet(tab_name)
            sheet.sheet_view.showGridLines = False
            sheet.sheet_view.zoomScale = 130
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.oddHeader.left.text = tab_name
            sheet.oddFooter.right.text = "Pairwise credited score  ·  row versus column"
            sheet.freeze_panes = "B2"

            header_row = 1
            corner = sheet.cell(header_row, 1, None)
            corner.fill = header_fill
            corner.border = thin
            for col_index, version in enumerate(versions, start=2):
                cell = sheet.cell(header_row, col_index, version)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = thin
            row_cursor = 2
            first_data = row_cursor
            numeric_cells = []
            magnitudes: list[float] = []
            for version in versions:
                label_cell = sheet.cell(row_cursor, 1, version)
                label_cell.font = label_font
                label_cell.fill = label_fill
                label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                label_cell.border = thin
                for col_index, other in enumerate(versions, start=2):
                    cell = sheet.cell(row_cursor, col_index, None)
                    cell.border = thin
                    cell.alignment = center
                    cell.font = score_font
                    if version == other:
                        cell.fill = diagonal_fill
                        cell.value = "—"
                        cell.font = Font(name="Calibri", size=11, color="B0B0B0")
                        continue
                    bucket = pairs.get((version, other))
                    if bucket is None:
                        continue
                    value = aggregate(bucket.values, how)
                    cell.value = value
                    cell.number_format = "0.##"
                    numeric_cells.append(cell.coordinate)
                    magnitudes.append(abs(value))
                    judge_bits = ", ".join(
                        f"{judge}: {score:g}" if judge else f"{score:g}"
                        for judge, score in zip(bucket.judges, bucket.values)
                    )
                    comment = Comment(
                        f"n={len(bucket.values)}\n{judge_bits}\n{agg_word}={value:g}",
                        "pairwise_scores",
                        width=220,
                        height=70,
                    )
                    cell.comment = comment
                sheet.row_dimensions[row_cursor].height = 20
                row_cursor += 1
            last_data = row_cursor - 1
            if numeric_cells and magnitudes:
                span = max(magnitudes) or 1
                score_range = f"B{first_data}:{get_column_letter(1 + len(versions))}{last_data}"
                sheet.conditional_formatting.add(
                    score_range,
                    ColorScaleRule(
                        start_type="num",
                        start_value=-span,
                        start_color="F8696B",
                        mid_type="num",
                        mid_value=0,
                        mid_color="FFFFFF",
                        end_type="num",
                        end_value=span,
                        end_color="63BE7B",
                    ),
                )
            sheet.row_dimensions[header_row].height = 22
            sheet.column_dimensions["A"].width = 14
            for col_index in range(2, 1 + len(versions) + 1):
                sheet.column_dimensions[get_column_letter(col_index)].width = 10
            sheet.freeze_panes = "B2"

    if not index_rows:
        raise SystemExit("no pairwise scores were found in the input")

    add_description_sheet(workbook, cubes, index_rows, how)
    if notes.get("truncated"):
        desc = workbook["Description"]
        # Find the "This file" row and append the truncation warning.
        for row in range(1, desc.max_row + 1):
            if desc.cell(row, 1).value == "This file":
                current = desc.cell(row, 2).value or ""
                desc.cell(row, 2).value = (
                    current
                    + " The input used here was a truncated head listing, so only those rows are included."
                )
                break

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def print_preview(cubes, how: str) -> None:
    for run in sorted(cubes):
        print(f"\n== {run}")
        rounds = cubes[run]

        def round_sort_key(label: str):
            return (0, int(label)) if str(label).isdigit() else (1, str(label))

        for round_id in sorted(rounds, key=round_sort_key):
            pairs = rounds[round_id]
            versions = sorted({a for pair in pairs for a in pair}, key=version_sort_key)
            title = "all rounds" if round_id == "all" else f"round {round_id}"
            print(f"  [{title}] {len(versions)} versions, {len(pairs)} directed pairs")
            header = " ".join(f"{v:>8}" for v in versions)
            print(f"  {'':>8} {header}")
            for version in versions:
                cells = []
                for other in versions:
                    if version == other:
                        cells.append(f"{'—':>8}")
                        continue
                    bucket = pairs.get((version, other))
                    if bucket is None:
                        cells.append(f"{'':>8}")
                    else:
                        cells.append(f"{aggregate(bucket.values, how):>8.2f}")
                print(f"  {version:>8} {''.join(cells)}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "inputs",
        nargs="+",
        type=Path,
        help="run directory, CSV file(s), or a head-dump text file",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("pairwise_scores.xlsx"),
        help="output workbook (default: pairwise_scores.xlsx)",
    )
    parser.add_argument(
        "--agg",
        choices=("mean", "median", "sum"),
        default="mean",
        help="how to combine multiple judge rows for the same ordered pair (default: mean)",
    )
    parser.add_argument(
        "--pool-rounds",
        action="store_true",
        help="average every round into a single matrix (default: one matrix per round)",
    )
    parser.add_argument(
        "--all-directions",
        action="store_true",
        help=("keep raw non-own frame rows for audit; `received` rows are "
              "negations of judgments already counted and can double count"),
    )
    return parser


def main(argv: list[str] | None = None) -> None:
    args = build_parser().parse_args(argv)
    cubes, notes = collect(
        args.inputs,
        how=args.agg,
        pool_rounds=args.pool_rounds,
        own_only=not args.all_directions,
    )
    if notes["directions"]:
        print("directions:", ", ".join(notes["directions"]))
    if notes["skipped_direction"]:
        print(f"skipped {notes['skipped_direction']} non-own rows")
    if notes["skipped_incomplete"]:
        print(f"skipped {notes['skipped_incomplete']} incomplete/self/non-numeric row(s)")
    if notes["truncated"]:
        print("warning: input is a truncated head listing; scores cover only those lines")
    print_preview(cubes, args.agg)
    write_workbook(cubes, notes, args.output, args.agg)
    sheets = sum(len(rounds) for rounds in cubes.values())
    print(f"\nwrote {args.output.resolve()} "
          f"({sheets} score sheet{'s' if sheets != 1 else ''})")


if __name__ == "__main__":
    main()
