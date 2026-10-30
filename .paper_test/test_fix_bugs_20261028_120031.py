#!/usr/bin/env python3
"""Regression checks for the 2026-10-28 12:00 fix-bugs wave (F01..F80 ids).

Run:  python3 .paper_test/test_fix_bugs_20261028_120031.py

Each check pins one finding the verify-findings stage kept as a TRUE_POSITIVE.
It fails on the pre-fix tree and passes on the repaired one.
"""
from __future__ import annotations

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import types
import zipfile
from pathlib import Path

WS = Path(__file__).resolve().parent.parent
TMPDIRS = []
FAILS = []


def load(name: str, rel: str):
    spec = importlib.util.spec_from_file_location(name, str(WS / rel))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


pp = load("fixbugs120031_pp", "paper_pipeline.py")
pf = load("fixbugs120031_pf", "paper_docx_format.py")
cc = load("fixbugs120031_cc", "paper-skills/paper-review/scripts/convert_corpus.py")


def scratch(prefix: str) -> Path:
    tmp = Path(tempfile.mkdtemp(prefix=prefix))
    TMPDIRS.append(tmp)
    return tmp


def check(name: str, cond, detail="") -> None:
    print(f"[{'ok ' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail else ""))
    if not cond:
        FAILS.append(name)


def big_docx(path: Path, size: int = 70 * 1024 * 1024) -> None:
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
           '2006/main"><w:body><w:p><w:r><w:t>' + ("x" * size)
           + '</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", xml)


LIMITS = {"abstract": {"cap": None, "relaxation": None},
          "main text": {"cap": None, "relaxation": None}, "cover letter": {}}
COUNTING = {"main_text": {"excludes": ["abstract", "methods", "references",
                                       "figure-legends"]}}


def test_f01_f10() -> None:
    print("\n== F01/F10: reserved-temp and append TOCTOU ==")
    root = scratch("f01_")
    victim = root / "victim"
    victim.write_text("ORIG", encoding="utf-8")
    dst = root / "out"
    tmp = pp.tmp_path_for(dst)
    os.unlink(tmp)
    os.symlink(victim, tmp)
    try:
        pp._atomic_write_bytes(tmp, dst, lambda f: f.write("ATTACK"))
        raised = False
    except OSError:
        raised = True
    check("F01 a swapped temp symlink is refused, not written through",
          raised and victim.read_text() == "ORIG")

    log = root / "log"
    os.symlink(victim, log)
    pp.append_text_no_follow(log, "APPENDED")
    check("F10 append replaces a planted symlink instead of following it",
          victim.read_text() == "ORIG" and not log.is_symlink()
          and log.read_text() == "APPENDED")


def test_f02_f08_f09_f60() -> None:
    print("\n== F02/F08/F09/F60: redline and template writers ==")
    try:
        from docx import Document
    except Exception as exc:                                    # noqa: BLE001
        check("python-docx is available for the DOCX checks", False, str(exc))
        return
    root = scratch("f02_")
    base, rev, out = root / "b.docx", root / "r.docx", root / "o.docx"
    Document().save(base)
    Document().save(rev)
    victim = root / "victim"
    victim.write_text("ORIG", encoding="utf-8")
    os.symlink(victim, out.with_suffix(".docx.tmp"))
    pp.builtin_tracked_changes(base, rev, out)
    check("F02 the fixed '<out>.docx.tmp' symlink is not used",
          victim.read_text() == "ORIG" and not out.is_symlink())

    check("F08 OUT == base is refused before the input is unlinked",
          pp.redline_one_pair("builtin", [], base, rev, base)
          and base.exists())

    tpl = root / "tpl.docx"
    Document().save(tpl)
    out2 = root / "templated.docx"
    victim2 = root / "victim2"
    victim2.write_text("ORIG", encoding="utf-8")
    os.symlink(victim2, out2)
    res = pf.apply_word_template(base, out2, tpl)
    check("F09 apply_word_template replaces a symlink atomically",
          res.get("ok") and victim2.read_text() == "ORIG" and not out2.is_symlink())

    # R05: the PKG-2 writers used to reopen the reserved temp by pathname.
    pkg_src = root / "pkg2_src.docx"
    doc = Document()
    doc.add_paragraph("This paragraph is kept.")
    doc.add_paragraph("The exact command is in code/analysis.py.")
    doc.save(pkg_src)
    pkg_out = root / "pkg2_out.docx"
    tmp_out = pp.tmp_path_for(pkg_out)       # reserve a sibling, as p12 does
    os.unlink(tmp_out)
    victim3 = root / "victim3"
    victim3.write_text("ORIG", encoding="utf-8")
    os.symlink(victim3, tmp_out)

    def finder(text):
        needle = "code/analysis.py"
        i = text.find(needle)
        return [(i, i + len(needle))] if i >= 0 else []

    res = pf.strip_reference_sentences(pkg_src, tmp_out, finder)
    check("F02 the package-2 rewrite does not write through a planted temp symlink",
          res.get("ok") and victim3.read_text() == "ORIG" and not tmp_out.is_symlink(),
          str(res)[:240])

    root2 = scratch("f60_")

    def make(p, rsid):
        xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/'
               'wordprocessingml/2006/main"><w:body><w:p w:rsidR="' + rsid
               + '"><w:pPr><w:spacing w:after="0"/></w:pPr><w:r><w:t>Same text.'
                 '</w:t></w:r></w:p></w:body></w:document>')
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", xml)

    b2, r2, o2 = root2 / "b.docx", root2 / "r.docx", root2 / "o.docx"
    make(b2, "00000001")
    make(r2, "00000002")
    info = pp.builtin_tracked_changes(b2, r2, o2)
    check("F60 a w:rsid-only resave is not a whole-paragraph change",
          info["paragraphs_inserted"] == 0 and info["paragraphs_deleted"] == 0,
          str(info))


def test_f03_f04_f05_f06() -> None:
    print("\n== F03/F04/F05/F06: untrusted-input budgets ==")
    root = scratch("f03_")
    fifo = root / "state.json"
    os.mkfifo(fifo)
    state, problem = pp.load_state_file(fifo)
    check("F03 a FIFO at state.json is refused, not opened", state is None
          and "regular file" in problem, problem)

    big = root / "big.docx"
    big_docx(big)
    text, notes = cc.docx_to_text(str(big))
    check("F04 the skill converter refuses the 70 MiB part", text == ""
          and any("budget" in n for n in notes), notes[:2])
    check("F05 the pipeline reader refuses the 70 MiB part",
          pp._docx_paragraphs(big) is None)
    try:
        pp._docx_root(big)
        raised = False
    except OSError:
        raised = True
    check("F05 _docx_root refuses the part instead of reading it", raised)

    sb = scratch("f06_")
    (sb / "revision_audit").mkdir()
    (sb / "revised").mkdir()
    plan = sb / "revision_audit" / "REPAIR_PLAN.json"
    plan.write_text('{"rows": [{"id": "RP-1", "file": "main.docx", "pad": "'
                    + ("x" * (65 * 1024 * 1024)) + '"}]}', encoding="utf-8")
    check("F06 an over-cap REPAIR_PLAN is not parsed",
          pp._revision_plan_rows(sb) is None)
    check("F06 read_text_bounded honours the cap",
          pp.read_text_bounded(plan) is None)


def test_f13_f14() -> None:
    print("\n== F13/F14: audit is disposition-only; repair plan is described ==")
    sb = scratch("f13_")
    (sb / "guided_review" / "round2").mkdir(parents=True)
    (sb / "guided_review" / "round2" / "probe_results.md").write_text(
        "# probe\n\n| probe id | disposition |\n|---|---|\n"
        "| P2-001 | executed - no anomalies - checked: base/x |\n", encoding="utf-8")
    (sb / "audit").mkdir()
    (sb / "audit" / "PROBE_AUDIT.md").write_text(
        "| probe id | disposition | verdict | evidence |\n|---|---|---|---|\n"
        "| P2-001 | no anomalies | promoted to AU-001 | x |\n", encoding="utf-8")
    (sb / "audit" / "audit.json").write_text('{"adds": []}', encoding="utf-8")
    errs = pp.probe_audit_problems(sb)
    check("F13 a probe promotion is refused with the UNRESOLVED route",
          any("retired" in e and "UNRESOLVED.md" in e for e in errs), errs[:2])
    (sb / "audit" / "PROBE_AUDIT.md").write_text(
        "| probe id | disposition | verdict | evidence |\n|---|---|---|---|\n"
        "| P2-001 | no anomalies | stands | x |\n", encoding="utf-8")
    check("F13 a disposition-only probe audit passes",
          pp.probe_audit_problems(sb) == [])
    check("F13 audit.json still rejects any AU-* add",
          bool(pp.audit_artifact_problems(
              {"dispositions": [{"id": "F-001", "verdict": "confirm"}],
               "adds": [{"id": "AU-1"}]}, ["F-001"])))

    plan_sb = scratch("f14_")
    (plan_sb / "revision_audit").mkdir()
    (plan_sb / "revised").mkdir()
    (plan_sb / "revision_audit" / "REPAIR_PLAN.json").write_text(
        json.dumps({"rows": [{"id": "RP-001", "file": "main.docx"}]}), encoding="utf-8")
    check("F14 a plan row with no described repair is rejected",
          pp._revision_plan_rows(plan_sb) is None)
    (plan_sb / "revision_audit" / "REPAIR_PLAN.json").write_text(json.dumps(
        {"rows": [{"id": "RP-001", "file": "main.docx",
                   "location": "Results p3", "wrong": "n = 12",
                   "correct": "n = 13", "basis": "raw_data/t.csv",
                   "edit": "replace n = 12", "evidence": "F-001"}]}),
        encoding="utf-8")
    check("F14 a fully described row is accepted",
          pp._revision_plan_rows(plan_sb) is not None)
    (plan_sb / "revised" / "REPAIR_LEDGER.json").write_text(json.dumps(
        {"rows": [{"id": "RP-001", "verdict": "confirmed", "reason": "applied"}]}),
        encoding="utf-8")
    described = pp._revision_plan_rows(plan_sb)
    check("F14 a confirmed ledger row must quote the applied text",
          bool(pp._revision_repair_ledger(plan_sb, described)[0]))
    pre, cur = scratch("f14pre_"), scratch("f14cur_")
    (pre / "main.md").write_text("Results: n = 12 in the cohort.", encoding="utf-8")
    (cur / "main.md").write_text("THE WHOLE FILE WAS REWRITTEN", encoding="utf-8")
    (cur / "REPAIR_LEDGER.json").write_text(json.dumps(
        {"rows": [{"id": "RP-001", "verdict": "confirmed", "reason": "applied",
                   "before": "n = 12", "after": "n = 13"}]}), encoding="utf-8")
    rows = [dict(described[0], file="main.md")] if described else []
    check("F14 a whole-file rewrite outside the quoted span is a violation",
          pp._revised_repair_violations(pre, cur, rows, {"RP-001"}) == ["main.md"])


def test_f26_f29_f31() -> None:
    print("\n== F26/F29/F31: document set, census opportunities, scope gate ==")
    b = scratch("f26b_")
    c = scratch("f26c_")
    (b / "figure-a.docx").write_text("A", encoding="utf-8")
    (c / "figure-b.docx").write_text("B", encoding="utf-8")
    res = pp.document_set_check([(b, "", ())], [(c, "", ())])
    check("F26 a single-letter VERSION-token rename is still the same document",
          res["missing"] == [] and res["added"] == [], str(res))
    b2, c2 = scratch("f26b2_"), scratch("f26c2_")
    (b2 / "Main.docx").write_text("same", encoding="utf-8")
    (c2 / "main.docx").write_text("same", encoding="utf-8")
    res = pp.document_set_check([(b2, "", ())], [(c2, "", ())])
    check("F26 a case-only name change is reported, not silently paired",
          res["content_preserved_under_other_name"] == ["Main.docx"], str(res))
    b3, c3 = scratch("f26b3_"), scratch("f26c3_")
    (b3 / "Main.docx").write_text("ONE", encoding="utf-8")
    (c3 / "main.docx").write_text("TWO", encoding="utf-8")
    res = pp.document_set_check([(b3, "", ())], [(c3, "", ())])
    check("F26 a case-only rename with edited bytes is lost + added",
          res["missing"] == ["Main.docx"] and res["added"] == ["main.docx"],
          str(res))

    comp = {"opponent_label": "v1", "score": 1, "reason": "x", "resolved": [],
            "introduced": [{"check": "M1", "tier": "consistency",
                            "severity": "minor",
                            "evidence": "line 1: abc def ghi"}], "checks": {}}
    obs = [("j1", "a1", "w1", comp), ("j1", "a1", "w2", comp)]
    census = pp.build_issue_census(obs, ["a1", "w1", "w2"], sessions_expected=2)
    check("F29 own_opps counts opponent comparisons, own_sessions counts sessions",
          census["a1"]["own_opps"] == 2 and census["a1"]["own_sessions"] == 1,
          str({k: census["a1"][k] for k in ("own_opps", "own_sessions")}))

    sb = scratch("f31_")
    (sb / "base").mkdir(parents=True)
    (sb / "guided_review" / "artifacts").mkdir(parents=True)
    (sb / "prior_round").mkdir()
    (sb / "guided_review" / "findings.md").write_text("#\n", encoding="utf-8")
    wanted = list(pp.REQUIRED_REVIEW_CHECKS) + [
        "M18", "M19", "M20", "M21", "M22", "M23", "M24", "M25", "M26", "M27",
        "M28", "M29", "J5", "M30", "M31", "M32", "M33", "M34", "M35", "M36"]
    coverage = [{"check": c, "disposition": "clean -- basis: a.md"} for c in wanted]
    coverage.append({"check": "M37", "disposition": "out of scope -- proposed check"})
    for name in ("M1.md", "M27_evidence_coverage.md", "M34_source_data_coverage.md"):
        (sb / "guided_review" / "artifacts" / name).write_text(
            "|check|d|\n|---|---|\n|x|clean|", encoding="utf-8")
    errs, warns = [], []
    pp.check_review_contract(types.SimpleNamespace(cfg={}), sb,
                             {"submission_dir": "./base", "findings": [],
                              "coverage": coverage}, errs, warns, scope="full")
    check("F31 an extra M37 row marked out of scope is not an in-scope error",
          not any("M37" in e for e in errs), [e for e in errs if "M37" in e])


def test_f34_f38() -> None:
    print("\n== F34-F38: M19 length scan ==")
    rows = pp._length_rows_for_lines(
        ["Abstract", "This is the abstract text.", "Keywords: alpha, beta",
         "This is the first body paragraph.", "Second body paragraph."],
        "x.docx", limits=LIMITS, counting=COUNTING)
    main = [r for r in rows if r["section"] == "main text"]
    check("F34 a body that follows Keywords directly is counted",
          main and main[0]["words"] == 9, rows)
    comma_body = "We studied cells, mice, and rats."
    rows = pp._length_rows_for_lines(
        ["Abstract", "This is the abstract text.", "Keywords: alpha, beta",
         comma_body, "Second body paragraph."],
        "x.docx", limits=LIMITS, counting=COUNTING)
    main = [r for r in rows if r["section"] == "main text"]
    check("F34 a comma-bearing body sentence is not a keyword continuation",
          main and main[0]["words"] == 9
          and not pp._length_keyword_continuation(comma_body), rows)

    rows = pp._length_rows_for_lines(
        ["Title", "Introduction", "Body text.", "Summary", "We conclude."],
        "x.docx", limits=LIMITS, counting=COUNTING)
    check("F35 a concluding Summary is not a second abstract",
          not any(r["section"] == "abstract" for r in rows)
          and any(r["section"] == "main text" and r["words"] >= 2 for r in rows),
          rows)

    root = scratch("f36_")
    (root / "main.tex").write_text(
        "\\documentclass{article}\\begin{document}\\begin{abstract}Short abstract."
        "\\end{abstract}\\section{Introduction}\\input{intro}\\end{document}\n",
        encoding="utf-8")
    (root / "intro.tex").write_text(
        "This is the body text with ten words in it total.\n", encoding="utf-8")
    info = pp.scan_lengths_in_sources([(root, "", ())],
                                      pp.load_venue_profile("generic"))
    main = [r for r in info["rows"] if r["document"] == "main.tex"
            and r["section"] == "main text"]
    check("F36 an \\input child is counted in its master and not twice",
          main and main[0]["words"] == 12
          and "intro.tex" in info["skipped"], (info["rows"], info["skipped"]))

    root_c = scratch("f36c_")
    (root_c / "main.tex").write_text(
        "\\documentclass{article}\\begin{document}\\begin{abstract}Short abstract."
        "\\end{abstract}\\section{Introduction}\nReal body words here.\n"
        "% \\input{old}\n\\end{document}\n", encoding="utf-8")
    (root_c / "old.tex").write_text(
        "\\section{Introduction}\nOld body words here.\n", encoding="utf-8")
    info = pp.scan_lengths_in_sources([(root_c, "", ())],
                                      pp.load_venue_profile("generic"))
    main = [r for r in info["rows"] if r["document"] == "main.tex"
            and r["section"] == "main text"]
    old = [r for r in info["rows"] if r["document"] == "old.tex"
           and r["section"] == "main text"]
    check("F36 a commented-out \\input is not expanded and its child keeps its row",
          main and main[0]["words"] == 4 and old and old[0]["words"] == 4,
          (info["rows"], info["skipped"]))

    lines, _, _ = pp._tex_length_scan(
        "\\section{Introduction}\nBody text here.\n\\printbibliography\n"
        "Ref one.\nRef two.\n")
    rows = pp._length_rows_for_lines(lines, "x.tex", limits=LIMITS,
                                     counting=COUNTING)
    check("F37 \\printbibliography ends the main-text count",
          rows and rows[0]["words"] == 3, rows)

    rows = pp._length_rows_for_lines(
        ["**Abstract**", "We studied cells.", "**Keywords:** a, b",
         "**Introduction**", "Body text here."],
        "x.md", limits=LIMITS, counting=COUNTING)
    check("F38 Markdown emphasis headings are recognised",
          len(rows) == 2 and rows[0]["words"] == 3 and rows[1]["words"] == 3, rows)


def test_f41_f44() -> None:
    print("\n== F41-F44: numbers ledger, table proof, caption grammar, conflicts ==")
    rows = pf.word_number_rows("a hundred samples")
    check("F41 'a hundred' is a count", bool(rows) and rows[0]["number"] == "100",
          rows)
    idx = pf.table_value_index([("x.csv", "label,rate\nA,50%\nB,20%\nC,30%\n")])
    check("F42 a percentage cell proves no plain 50",
          idx.get("50") is None, idx)
    tables = [("t.tsv", "name\tvalue\n" + "".join(
        f"{c}\t{v}\n" for c, v in [("a", 11), ("b", 12), ("c", 13),
                                    ("d", ""), ("e", ""), ("f", ""),
                                    ("g", ""), ("h", ""), ("i", ""), ("j", "")]))]
    idx = pf.table_value_index(tables)
    check("F42 the row count stays the reported every-row count (G-041)",
          idx.get("10") == "number of data rows in t.tsv", idx.get("10"))
    rows = [{"number": "10", "source": ""}, {"number": "36", "source": ""}]
    out = pf.reconcile_number_rows(rows, tables)
    check("F42 a row count is a reported hint, never a provenance proof",
          rows[0]["source"] == "" and rows[1]["source"].startswith("recomputed:"),
          out)
    check("F43 Table I is a display caption",
          bool(pf.TABLE_CAPTION_RE.match("Table I. Title")))
    check("F43 Table D uses the pipeline's roman grammar too",
          bool(pf.TABLE_CAPTION_RE.match("Table D. Title"))
          and bool(pf.TABLE_CAPTION_BARE_RE.match("Table D")))
    check("F43 a body sentence opening 'Figure C.' is not a caption",
          not pf.FIGURE_CAPTION_RE.match("Figure C. We next tested the assay in mice."))
    def op(j, e):
        return {"judge_run": j, "target": "a1", "opponent": "w1", "score": 1,
                "reason": "n samples",
                "resolved": [{"check": "M4", "tier": "correctness",
                              "severity": "minor", "evidence": e}],
                "introduced": [], "checks": {}}
    conflicts = pp.detect_round_judge_conflicts(
        [op("j1", "n=173 samples"), op("j2", "n=508 samples")])
    check("F44 a singleton-vs-singleton number disagreement is a conflict",
          any(c.get("kind") == "numeric_contradiction" for c in conflicts),
          conflicts)


def test_f77_f78_f80() -> None:
    print("\n== F77/F78/F80: one classifier, profile containment, dependency message ==")
    names = ("cover-letter-response.docx", "response_to_reviewers.docx",
             "cover_letter.docx")
    pairs = [(pp._is_cover_letter_lines([], n), pf.is_cover_letter_doc(n)) for n in names]
    check("F77 the pipeline and formatter cover-letter predicates agree",
          pairs == [(False, False), (False, False), (True, True)], pairs)

    node = shutil.which("node")
    if node:
        root = scratch("f78_")
        (root / "profiles" / "ok").mkdir(parents=True)
        (root / "profiles" / "ok" / "cordis.patch.yml").write_text("# patch\n",
                                                                   encoding="utf-8")
        proc = subprocess.run(
            [node, str(WS / "scripts" / "install-codex-mcp.mjs"),
             "--dsh-home", str(root), "--profile", "../../outside", "--dry-run"],
            capture_output=True, text=True)
        check("F78 a path-traversal --profile is rejected",
              proc.returncode != 0 and "invalid --profile" in proc.stderr,
              proc.stderr[-200:])
    else:
        check("F78 node is available for the installer check", False, "node not on PATH")

    script = (WS / "docxcompare.sh").read_text(encoding="utf-8", errors="replace")
    check("F80 a missing python3 is reported as a dependency error, output removed",
          "command -v python3" in script
          and "required to verify the comparison document" in script
          and 'rm -f -- "$OUTPUT"' in script)


def main() -> int:
    try:
        test_f01_f10()
        test_f02_f08_f09_f60()
        test_f03_f04_f05_f06()
        test_f13_f14()
        test_f26_f29_f31()
        test_f34_f38()
        test_f41_f44()
        test_f77_f78_f80()
    finally:
        for tmp in TMPDIRS:
            shutil.rmtree(tmp, ignore_errors=True)
    if FAILS:
        print(f"\n{len(FAILS)} fix-bugs check(s) FAILED")
        for name in FAILS:
            print("  - " + name)
        return 1
    print("\nALL FIX-BUGS CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
