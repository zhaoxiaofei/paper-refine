#!/usr/bin/env python3
"""PKG-6/7/8/9: statements a package makes about ITSELF that contradict each other.

The operator's result-driven audit of the two delivered packages
(`audit_data/2026-1015-2208-result-driven-audit-of-both-manuscripts/`) found
three defects that every existing rule was blind to, because each rule compares
one surface against ITSELF (a legend against its own figure, a reference against
its own entry) or against the raw data -- never two statements the same
submission makes:

  * `PKG-6` -- both packages' front matter said the manuscript was drafted
    "without the use of generative AI" and, in the same section, described
    running a generative-AI revision workflow over it. A journal that requires
    AI disclosure reads the denial as a false statement.
  * `PKG-7` -- the CNB Reporting Summary printed the scRNA commit as
    `c30bc1ab52f083b5761e31bec6888d7f2fd` (35 hex) while the main text carried
    the full 40: resolving the short form yields no object.
  * `PKG-8` -- the supplementary captions said "pairwise comparisons used
    606-1,989 paired cells" while Table S2's own note defined n as "the number
    of donors whose paired differences ... enter the test" (n = 9).

This suite pins the three detectors on hand-built positives, on the NEGATIVE
controls each one produced on the real corpora (scoped AI denials, citation
ranges, decimal versions), the shared document-text reader that silently
skipped the 4.6 MB CNB main text while it had a single global 4 MB cap, and the
two delivered packages themselves -- those last checks fail until the shipped
prose is corrected, which is how the defects were confirmed.

Run:  python3 .paper_test/test_contradictions_2026_1015b.py
"""
from __future__ import annotations

import importlib.util
import os
import sys
from pathlib import Path

WS = Path(os.environ.get("PAPER_WS") or Path(__file__).resolve().parent.parent)
spec = importlib.util.spec_from_file_location("paper_contr", str(WS / "paper_pipeline.py"))
pp = importlib.util.module_from_spec(spec)
sys.modules["paper_contr"] = pp
spec.loader.exec_module(pp)

FAILS = []
BASE = Path("/home/zhaoxiaofei/manuscript-postgrad-data")


def check(name, cond, detail=""):
    print(f"[{'ok ' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail else ""))
    if not cond:
        FAILS.append(name)


# The two shipped denial sentences (quoted from the delivered bytes), each
# followed by the AI-workflow sentence that contradicts it.
CNB_DISCLOSURE = (
    "Generative AI use\n"
    "X.F.Z. drafted this manuscript without generative AI, except for using GLM-5.2 and "
    "GLM-5.3 to recommend journals. After the draft was complete, X.F.Z. designed and "
    "developed a generative AI-based revision workflow "
    "(https://github.com/zhaoxiaofei/paper-refine) and used its default settings, which call "
    "the Codex interface (with DeepSeek V4.1 Flash behind the interface), to refine the "
    "manuscript.\n"
)
NH_DISCLOSURE = (
    "AI use: X.Zhao drafted this manuscript without the use of generative AI, with the sole "
    "exception of employing GLM-5.2 and GLM-5.3 to recommend journals. After the draft was "
    "completed, X.Zhao designed and developed a generative AI-based revision workflow "
    "(https://github.com/zhaoxiaofei/paper-refine) and applied its default settings which "
    "interface with Codex and use DeepSeek V4.1 Flash as the underlying model to refine the "
    "manuscript.\n"
)
# Scoped denials: true statements that must NOT be paired with an AI-writing
# disclosure. These are the negative controls the real corpora produced.
SCOPED_DENIALS = [
    "No AI tools were used for the statistical analysis; the text was refined with an "
    "AI-based revision workflow.",
    "We did not use automated writing tools for the figures. The text was refined with "
    "ChatGPT.",
    "All figures were generated without AI assistance; the manuscript text was refined using "
    "an AI-based revision workflow.",
]


def test_pkg6_ai_disclosure():
    for label, text in (("cnb", CNB_DISCLOSURE), ("neohet", NH_DISCLOSURE)):
        rows = pp.ai_disclosure_contradiction_rows([(f"{label}.docx", text)])
        check(f"PKG-6 fires on the shipped {label} disclosure", len(rows) == 1,
              str(rows[:1])[:120])
        if rows:
            check(f"PKG-6 names both sides in the {label} evidence",
                  "without" in rows[0]["evidence"]
                  and "generative AI-based revision" in rows[0]["evidence"],
                  rows[0]["evidence"][:140])
    for text in SCOPED_DENIALS:
        rows = pp.ai_disclosure_contradiction_rows([("x.docx", text)])
        check(f"PKG-6 stays quiet on a SCOPED denial ({text[:38]!r}...)", not rows,
              str(rows[:1]))
    # One statement per document: a denial in the cover letter and a workflow
    # description in the Methods are two different texts, not a contradiction.
    rows = pp.ai_disclosure_contradiction_rows(
        [("cover.docx", "This manuscript was drafted without generative AI."),
         ("main.docx", "We used a generative AI-based revision workflow on the text.")])
    check("PKG-6 does not pair statements from two different documents", not rows,
          str(rows[:1]))
    check("PKG-6 stays quiet on a denial alone",
          not pp.ai_disclosure_contradiction_rows(
              [("x.docx", "The authors drafted this manuscript without generative AI.")]))


FULL_SHA = "c30bc1ab52f083b5761e31bec6888d7f2fdc9ddb"
SHORT_SHA = FULL_SHA[:35]


def test_pkg7_truncated_identifier():
    rows = pp.truncated_identifier_rows(
        [("reporting-summary.pdf", f"https://github.com/x/commit/{SHORT_SHA} The study used "),
         ("mainText.docx", f"https://github.com/x/tree/{FULL_SHA}")])
    check("PKG-7 fires on the shipped truncated commit", len(rows) == 1,
          str(rows[:1])[:140])
    if rows:
        check("PKG-7 quotes the length of each spelling",
              "35 hex" in rows[0]["evidence"] and "40-hex" in rows[0]["evidence"],
              rows[0]["evidence"][:140])
    check("the two spellings in ONE document are not a contradiction",
          not pp.truncated_identifier_rows([("x.docx", f"{SHORT_SHA} and {FULL_SHA}")]))
    check("unrelated hex-like tokens are silent",
          not pp.truncated_identifier_rows(
              [("a.docx", "versions 1.22.0 and 2024.1 and "
                          "0b105331aa1fbbd94832ae308a84e133831fa807")]))
    check("a bare short token with no longer form is silent",
          not pp.truncated_identifier_rows([("x.docx", SHORT_SHA)]))


TABLE_NOTE = ("n is the effective sample size, that is, the number of donors whose paired "
              "differences of the two callers' per-donor medians enter the test.")
CAPTION = ("Sample sizes: n = 1,989 simulated cells from 9 donors; pairwise comparisons used "
           "606-1,989 paired cells depending on caller availability.")


def test_pkg8_unit_of_analysis():
    rows = pp.unit_of_analysis_rows([("supp.pdf", CAPTION), ("supp.pdf", TABLE_NOTE)])
    check("PKG-8 fires on the shipped caption/table mismatch", len(rows) == 1,
          str(rows[:1])[:140])
    if rows:
        check("PKG-8 names both units",
              "'cells'" in rows[0]["detail"] and "'donors'" in rows[0]["detail"],
              rows[0]["detail"][:160])
    check("PKG-8 stays quiet when the units agree",
          not pp.unit_of_analysis_rows(
              [("supp.pdf", CAPTION),
               ("supp.pdf", "n is the number of cells whose paired differences enter the test")]))
    check("PKG-8 stays quiet when no caption claims a sample size",
          not pp.unit_of_analysis_rows(
              [("supp.pdf", "Pairwise comparisons were used for the CN metrics."),
               ("supp.pdf", TABLE_NOTE)]))
    check("repeated captions produce ONE row per unit pair",
          len(pp.unit_of_analysis_rows([("supp.pdf", CAPTION * 5), ("supp.pdf", TABLE_NOTE)]))
          == 1)


def test_document_text_reader():
    """The reader that silently skipped a 4.6 MB manuscript (and figure exports)."""
    if not BASE.is_dir():
        print("skip: the operator's manuscript tree is not readable here")
        return
    root = (BASE / "cnb-30to31-1015-1812-669237d"
            / "final_clean_version_cnb-30to31-1015-1812-669237d")
    if not root.is_dir():
        print("skip: the CNB delivered package is not present")
        return
    texts = dict(pp.package_document_texts(root))
    main = [rel for rel in texts if "mainText" in rel and rel.endswith(".docx")]
    check("the document reader reaches a 4.6 MB manuscript docx", bool(main),
          str(list(texts))[:200])
    check("the document reader skips per-figure source-data exports",
          not [rel for rel in texts if rel.startswith("supplementary_source_files/")],
          str(list(texts))[:200])
    check("every carried document is read whole (no silent truncation at a byte cap)",
          bool(main) and len(texts[main[0]]) > 40_000,
          str(len(texts.get(main[0], "")) if main else 0))


def test_evidence_and_readiness_wiring():
    rows = pp.truncated_identifier_rows([("rs.pdf", SHORT_SHA), ("main.docx", FULL_SHA)])
    rdy = pp.package_readiness({"hygiene": {"contradictions": rows}})
    check("a contradiction row BLOCKS readiness (the author decides which side holds)",
          bool(rows) and any(f.get("rule") == "PKG-7" and not f.get("advisory")
                             for f in rdy.get("findings") or []),
          str(rdy.get("findings"))[:160])


def test_pkg9_tool_definition():
    """PKG-9: a table's affinity claim for a tool the prose never says it of.

    The shipped row is quoted from the delivered neohet main text; the control
    is the same table's NetMHCpan row, whose prose DOES name the tool beside
    the quantity ("NetMHCpan-predicted half-maximal inhibitory concentration
    (IC50)").
    """
    import tempfile
    import shutil
    try:
        import docx                                              # noqa: PLC0415
    except Exception:                                            # noqa: BLE001
        print("skip: python-docx is not installed here")
        return
    d = Path(tempfile.mkdtemp(prefix="hyg_pkg9_"))
    try:
        def make(name, prose, rows):
            doc = docx.Document()
            for para in prose:
                doc.add_paragraph(para)
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        t.cell(i, j).text = cell
            doc.save(str(d / name))

        # The Methods sentence is written as a SENTENCE (it ends in a period):
        # a support test that runs to the paragraph break would read the next
        # paragraph as part of it, which is what the first cut of the rule did.
        good = ["To limit peak random-access memory, candidates whose NetMHCpan-predicted "
                "half-maximal inhibitory concentration (IC50) exceeds 10,000 nM are discarded.",
                "MMP_aff ranks candidates by the percentile rank of the MixMHCpred score "
                "against a length- and allotype-matched background of random peptides."]
        rows = [["MMP_aff", "PWM-derived", "MixMHCpred_aff", "Affinity (IC50) percentile",
                 "Percentile rank of the predicted IC50 of the pMHC pair (MixMHCpred)"],
                ["NMP_aff", "NN ensemble", "NetMHCpan_aff", "Affinity (IC50)",
                 "Predicted IC50 of the pMHC pair (NetMHCpan)"]]
        make("mixed.docx", good, rows)
        got = pp.affinity_definition_rows(d)
        tools = {r["evidence"].split("for ", 1)[1].split(",", 1)[0].strip("'") for r in got}
        check("PKG-9 flags the MixMHCpred IC50 claim", any("MixMHCpred" in t for t in tools),
              str(tools))
        check("PKG-9 does NOT flag the well-defined NetMHCpan row",
              not any("NetMHCpan" in t for t in tools), str(tools))
        check("PKG-9 quotes the row that carries the claim",
              got and "Affinity (IC50) percentile" in got[0]["evidence"], str(got[:1])[:160])
    finally:
        shutil.rmtree(d, ignore_errors=True)


def test_real_packages_have_no_open_contradiction():
    """The end-to-end half: the two delivered packages as they stand.

    After this round's fixes both packages are clean -- EXCEPT one authored row
    the audit filed and the author has yet to act on: the neohet Table 2
    "Affinity (IC50) percentile" claim for MixMHCpred (PKG-9), which the same
    round reports in the delivered status and `readiness.findings`. The check
    names it explicitly instead of asserting "no rows at all", so a NEW
    contradiction still fails the suite and a fixed table turns it green.
    """
    if not BASE.is_dir():
        print("skip: the operator's manuscript tree is not readable here")
        return
    expected_open = {"neohet-13to14-1015-1812-669237d":
                     {("PKG-9", "neohetero-15-2-mainText-Fron")}}
    for name in ("cnb-30to31-1015-1812-669237d", "neohet-13to14-1015-1812-669237d"):
        root = BASE / name / f"final_clean_version_{name}"
        if not root.is_dir():
            print(f"skip: {root} is not present")
            continue
        rows = pp.package_contradiction_rows(root)
        got = {(r["rule"], r["file"][:28]) for r in rows}
        check(f"{name}: only the AUDITED, author-owned contradiction is still open",
              got == expected_open.get(name, set()),
              f"open={sorted(got)}")


def main() -> int:
    test_pkg6_ai_disclosure()
    test_pkg7_truncated_identifier()
    test_pkg8_unit_of_analysis()
    test_document_text_reader()
    test_evidence_and_readiness_wiring()
    test_pkg9_tool_definition()
    test_real_packages_have_no_open_contradiction()
    if FAILS:
        print(f"\n{len(FAILS)} FAILED: " + "; ".join(FAILS))
        return 1
    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
