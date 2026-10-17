#!/usr/bin/env python3
"""Pairwise score-workbook checks: directed rows, provenance and workbook layout.

Run:  python3 .paper_test/test_pairwise_scores.py

The important invariant is that `paper_pipeline.py` writes BOTH frames of one
judge comparison into `round<N>_raw_scores.csv`: the target's `own` row and the
opponent's `received` negation. A directed matrix must use the own row only;
using the received row as well fabricates a judgment in the opposite direction.
"""
from __future__ import annotations

import csv
import importlib.util
import io
import os
import shutil
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

WS = Path(os.environ.get("PAPER_WS") or Path(__file__).resolve().parent.parent)
FAILS = []
TMPDIRS = []


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, str(path))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


ps = _load("paper_pairwise_scores_test", WS / "pairwise_scores.py")

FIELDS = [
    "round", "member", "credited", "direction", "opponent_id", "opponent_label",
    "source_judge_run", "source_judge_index", "sheet_target", "score_as_written",
    "basis", "resolved", "introduced", "reason", "member_n", "member_median",
    "member_mean",
]


def check(name, cond, detail=""):
    print(f"[{'ok ' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail else ""))
    if not cond:
        FAILS.append(name)


def scratch(prefix: str) -> Path:
    d = Path(tempfile.mkdtemp(prefix=prefix))
    TMPDIRS.append(d)
    return d


def cleanup():
    for d in TMPDIRS:
        shutil.rmtree(d, ignore_errors=True)


def score_row(round_id: int, member: str, opponent: str, score: float,
              direction: str, judge: str, index: int,
              target: str | None = None) -> dict:
    return {
        "round": round_id,
        "member": member,
        "credited": score,
        "direction": direction,
        "opponent_id": opponent,
        "opponent_label": "v1",
        "source_judge_run": judge,
        "source_judge_index": index,
        "sheet_target": target or member,
        "score_as_written": score,
        "basis": "correctness",
        "resolved": "1",
        "introduced": "0",
        "reason": "fixture",
        "member_n": 2,
        "member_median": score,
        "member_mean": score,
    }


def write_scores(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def test_parse_score():
    print("\n== finite numeric scores only ==")
    check("a normal credited score parses", ps.parse_score(" -2 ") == -2.0)
    check("blank is skipped", ps.parse_score("") is None)
    for value in ("nan", "NaN", "inf", "-inf", "1e309"):
        check(f"non-finite {value!r} is skipped", ps.parse_score(value) is None)


def test_directed_collection():
    print("\n== own rows only; never synthesize the opposite direction ==")
    run = scratch("pairwise_run_") / "cnb-12to13-demo"
    write_scores(run / "reports" / "round1_raw_scores.csv", [
        score_row(1, "a1", "a2", 2, "own", "j1", 1),
        score_row(1, "a1", "a2", 4, "own", "j2", 2),
        score_row(1, "a2", "a1", 1, "own", "j2", 2),
        # Derived mirrors of those two judgments. They must never enter the
        # directed matrix as if they were independent judgments.
        score_row(1, "a2", "a1", -2, "received", "j1", 1, target="a1"),
        score_row(1, "a2", "a1", -4, "received", "j2", 2, target="a1"),
        score_row(1, "a2", "a3", 1, "own", "j1", 1),
        score_row(1, "a3", "a1", 5, "mirror", "j1", 1),
        score_row(1, "a1", "a1", 3, "own", "j1", 1),
        score_row(1, "a1", "a3", "nan", "own", "j1", 1),
    ])
    write_scores(run / "reports" / "round2_raw_scores.csv", [
        score_row(2, "a1", "a2", -1, "own", "j1", 1),
    ])

    cubes, notes = ps.collect([run], how="mean", pool_rounds=False, own_only=True)
    pairs = cubes[run.name]["1"]
    check("the two own judgments are averaged",
          pairs[("a1", "a2")].values == [2.0, 4.0], str(pairs.get(("a1", "a2"))))
    check("a genuinely judged opposite direction survives unchanged",
          pairs[("a2", "a1")].values == [1.0], str(pairs.get(("a2", "a1"))))
    check("received rows do not contaminate that opposite direction",
          -2.0 not in pairs[("a2", "a1")].values
          and -4.0 not in pairs[("a2", "a1")].values,
          str(pairs.get(("a2", "a1"))))
    check("a genuinely judged second pair survives",
          pairs[("a2", "a3")].values == [1.0], str(pairs.get(("a2", "a3"))))
    check("non-own rows are counted as skipped",
          notes["skipped_direction"] == 3, str(notes))
    check("self and non-finite rows are skipped",
          notes["skipped_incomplete"] == 2, str(notes))
    check("round 2 stays a separate matrix",
          cubes[run.name]["2"][("a1", "a2")].values == [-1.0])

    all_cubes, _ = ps.collect([run], how="mean", pool_rounds=False, own_only=False)
    all_pairs = all_cubes[run.name]["1"]
    check("the audit mode retains the raw received rows",
          sorted(all_pairs[("a2", "a1")].values) == [-4.0, -2.0, 1.0],
          str(all_pairs.get(("a2", "a1"))))
    return run, cubes, notes


def test_head_dump():
    print("\n== head-listing blocks ==")
    root = scratch("pairwise_head_")
    virtual = root / "cnb-9to10-head" / "reports" / "round2_raw_scores.csv"
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=FIELDS)
    writer.writeheader()
    writer.writerow(score_row(2, "a2", "a1", 3, "own", "j1", 1))
    csv_text = out.getvalue()
    dump = root / "head.txt"
    dump.write_text(f"==> {virtual} <==\n{csv_text}", encoding="utf-8")
    cubes, notes = ps.collect([dump], how="mean", pool_rounds=False, own_only=True)
    check("a head dump is parsed into its own run and round",
          cubes[virtual.parent.parent.name]["2"][("a2", "a1")].values == [3.0],
          str(cubes))
    check("a head dump is marked truncated in the notes", notes["truncated"] is True)


def test_workbook():
    print("\n== workbook values and sheet count ==")
    run = scratch("pairwise_render_") / "cnb-12to13-demo"
    write_scores(run / "reports" / "round1_raw_scores.csv", [
        score_row(1, "a1", "a2", 2, "own", "j1", 1),
        score_row(1, "a1", "a2", 4, "own", "j2", 2),
        score_row(1, "a2", "a1", -2, "received", "j1", 1, target="a1"),
    ])
    write_scores(run / "reports" / "round2_raw_scores.csv", [
        score_row(2, "a1", "a2", -1, "own", "j1", 1),
    ])
    cubes, notes = ps.collect([run], how="mean", pool_rounds=False, own_only=True)
    out = scratch("pairwise_xlsx_") / "scores.xlsx"
    ps.write_workbook(cubes, notes, out, "mean")

    from openpyxl import load_workbook
    book = load_workbook(out, data_only=True)
    check("Description is the first sheet", book.sheetnames[0] == "Description")
    check("one score sheet is written per round",
          book.sheetnames == ["Description", "12to13-demo_r1", "12to13-demo_r2"],
          str(book.sheetnames))
    sheet = book["12to13-demo_r1"]
    headers = [c.value for c in sheet[1]]
    rows = {sheet.cell(r, 1).value: r for r in range(2, sheet.max_row + 1)}
    check("the own mean is the matrix value",
          sheet.cell(rows["a1"], headers.index("a2") + 1).value == 3.0,
          str([c.value for c in sheet[rows["a1"]]]))
    check("the unjudged opposite direction stays blank",
          sheet.cell(rows["a2"], headers.index("a1") + 1).value is None,
          str(sheet.cell(rows["a2"], headers.index("a1") + 1).value))

    preview = io.StringIO()
    with redirect_stdout(preview):
        ps.main([str(run), "-o", str(out)])
    check("the completion line counts score sheets, not run directories",
          "2 score sheets" in preview.getvalue(), preview.getvalue()[-200:])


def main():
    try:
        test_parse_score()
        test_directed_collection()
        test_head_dump()
        test_workbook()
    finally:
        cleanup()
    print()
    if FAILS:
        print(f"{len(FAILS)} CHECK(S) FAILED:")
        for name in FAILS:
            print(f"  - {name}")
        return 1
    print("PAIRWISE SCORE WORKBOOK CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
